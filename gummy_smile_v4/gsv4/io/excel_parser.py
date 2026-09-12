"""Parser for the clinical reference workbook and the calibration workbook.

Cell coding rules come from the observer's own notes inside the workbook (spec §4.2):

    integer >= 1000   -> /1000        Turkish Excel read ``3.207`` as a thousands number
    ``*0.718``        -> 0.718        sub-millimetre value, the ``*`` avoids losing the 0
    ``-`` variants    -> 0.0 mm       no visible gingiva on that tooth (clinical decision)
    integer 100..999  -> /1000, flagged ambiguous (probably a dropped leading zero)
    small number / numeric text -> as is
    negative          -> NaN
Every cell keeps its ``parse_kind`` so that sensitivity analyses can include/exclude
each rule. The E label written next to each cell is used as a consistency check.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np
import openpyxl
import pandas as pd

from gsv4.io.naming import normalize_excel_name
from gsv4.rules.thresholds import NO_VISIBLE_GINGIVA, UNCLASSIFIED, label_for_mm

PARSE_KINDS = [
    "div1000", "star_sub1mm", "dash_zero", "ambiguous_100_999", "plain_mm",
    "negative", "missing", "unparsed",
]

_LABEL_RE = re.compile(r"^([ET])\s*(\d)(?:\s*[-–—]\s*[ET]?\s*(\d))?$", re.IGNORECASE)
_SEX_MAP = {"K": "F", "KADIN": "F", "F": "F", "E": "M", "ERKEK": "M", "M": "M"}


@dataclass(frozen=True)
class ParsedCell:
    value_mm: float
    parse_kind: str
    raw: Any


def _to_float(s: str) -> Optional[float]:
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def _numeric_rule(x: float, raw: Any) -> ParsedCell:
    if x < 0:
        return ParsedCell(math.nan, "negative", raw)
    if x >= 1000:
        return ParsedCell(x / 1000.0, "div1000", raw)
    if 100 <= x < 1000:
        return ParsedCell(x / 1000.0, "ambiguous_100_999", raw)
    if x < 10:
        return ParsedCell(float(x), "plain_mm", raw)
    return ParsedCell(math.nan, "unparsed", raw)


def parse_measurement_cell(raw: Any) -> ParsedCell:
    """Apply the coding rules to one measurement cell."""
    if raw is None:
        return ParsedCell(math.nan, "missing", raw)
    if isinstance(raw, (int, float, np.integer, np.floating)) and not isinstance(raw, bool):
        if isinstance(raw, float) and math.isnan(raw):
            return ParsedCell(math.nan, "missing", raw)
        return _numeric_rule(float(raw), raw)
    s = str(raw).strip()
    if s == "":
        return ParsedCell(math.nan, "missing", raw)
    if s.startswith("*"):
        v = _to_float(s[1:].strip())
        return ParsedCell(v, "star_sub1mm", raw) if v is not None else ParsedCell(math.nan, "unparsed", raw)
    if s.startswith("-"):
        v = _to_float(s)
        if v is not None:                       # e.g. "-684"
            return _numeric_rule(v, raw)
        return ParsedCell(0.0, "dash_zero", raw)  # "-", "-(mesafe yok)", "- "
    v = _to_float(s)
    if v is not None:
        return _numeric_rule(v, raw)
    return ParsedCell(math.nan, "unparsed", raw)


def parse_label(raw: Any) -> Optional[str]:
    """Normalise an E/T label cell: ``E1``, ``E1-E2``, ``-`` or None."""
    if raw is None:
        return None
    s = str(raw).strip()
    if s == "":
        return None
    if s.startswith("-") and not any(ch.isdigit() for ch in s):
        return "-"
    m = _LABEL_RE.match(s)
    if not m:
        return s.upper()
    prefix = m.group(1).upper()
    first, second = m.group(2), m.group(3)
    return f"{prefix}{first}" if second is None else f"{prefix}{first}-{prefix}{second}"


def parse_age(raw: Any) -> Optional[float]:
    if raw is None:
        return None
    if isinstance(raw, (int, float, np.integer, np.floating)) and not isinstance(raw, bool):
        return float(raw) if 0 < float(raw) < 120 else None
    s = str(raw).strip()
    if re.fullmatch(r"\d{1,3}", s) and 0 < int(s) < 120:
        return float(s)
    return None


def parse_sex(raw: Any) -> Optional[str]:
    if raw is None:
        return None
    s = str(raw).strip().upper().replace("İ", "I")
    return _SEX_MAP.get(s)


def _rows(path: Path, sheet: str) -> List[tuple]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"Sheet {sheet!r} not found in {path}; sheets: {wb.sheetnames}")
    return list(wb[sheet].iter_rows(values_only=True))


def _cell(row: tuple, idx: int) -> Any:
    return row[idx] if idx < len(row) else None


def read_high_sheet(path: Path, excel_cfg: Dict[str, Any]) -> pd.DataFrame:
    """One record per Excel row of the high-smile-line sheet, wide format."""
    cols = excel_cfg["high_columns"]
    teeth = excel_cfg["teeth"]
    header_row = int(excel_cfg["high_header_row"])
    records: List[Dict[str, Any]] = []
    for ridx, row in enumerate(_rows(path, excel_cfg["high_sheet"])):
        if ridx <= header_row:
            continue
        raw_name = _cell(row, cols["name"])
        if raw_name is None or str(raw_name).strip() == "":
            continue
        name = normalize_excel_name(raw_name)
        rec: Dict[str, Any] = {
            "excel_row": ridx, "raw_name": str(raw_name).strip(), "key": name.key,
            "base": name.base, "dot": name.dot, "age_prefix": name.age_prefix,
        }
        values = []
        n_incons = 0
        n_labelled = 0
        for i, (tcol, ecol, tcol2) in enumerate(zip(cols["teeth"], cols["etiology"], cols["treatment"]), start=1):
            cell = parse_measurement_cell(_cell(row, tcol))
            label = parse_label(_cell(row, ecol))
            expected = label_for_mm(cell.value_mm)
            if expected == NO_VISIBLE_GINGIVA:
                expected = "-"
            elif expected == UNCLASSIFIED:
                expected = None
            inconsistent = label is not None and expected is not None and label != expected
            rec[f"mm_{i}"] = cell.value_mm
            rec[f"kind_{i}"] = cell.parse_kind
            rec[f"raw_{i}"] = cell.raw
            rec[f"label_{i}"] = label
            rec[f"expected_label_{i}"] = expected
            rec[f"treatment_{i}"] = parse_label(_cell(row, tcol2))
            rec[f"label_inconsistent_{i}"] = bool(inconsistent)
            rec[f"tooth_{i}"] = teeth[i - 1]
            n_incons += int(inconsistent)
            n_labelled += int(label is not None)
            values.append(cell.value_mm)
        valid = [v for v in values if not math.isnan(v)]
        kinds = [rec[f"kind_{i}"] for i in range(1, 7)]
        rec["n_valid"] = len(valid)
        rec["complete"] = len(valid) == 6
        rec["mean_mm"] = float(np.mean(valid)) if len(valid) == 6 else math.nan
        rec["has_dash_zero"] = "dash_zero" in kinds
        rec["has_ambiguous_100_999"] = "ambiguous_100_999" in kinds
        rec["label_inconsistent_count"] = n_incons
        rec["n_labelled"] = n_labelled
        rec["age_raw"] = _cell(row, cols["age"])
        rec["sex_raw"] = _cell(row, cols["sex"])
        age = parse_age(rec["age_raw"])
        rec["age_source"] = "sheet" if age is not None else ("prefix" if name.age_prefix else None)
        rec["age"] = age if age is not None else (float(name.age_prefix) if name.age_prefix else math.nan)
        rec["prefix_age_matches_sheet"] = (name.age_prefix is not None and age is not None and float(name.age_prefix) == age)
        rec["sex"] = parse_sex(rec["sex_raw"])
        records.append(rec)
    df = pd.DataFrame(records)
    df["reference_label"] = df["mean_mm"].map(label_for_mm)
    return df


def parse_summary(df: pd.DataFrame) -> Dict[str, Any]:
    """Counts reported in parse_report.md."""
    kinds = pd.Series([df[f"kind_{i}"] for i in range(1, 7)]).explode() if False else pd.concat([df[f"kind_{i}"] for i in range(1, 7)])
    kind_counts = {k: int((kinds == k).sum()) for k in PARSE_KINDS}
    return {
        "n_rows": int(len(df)),
        "n_complete": int(df["complete"].sum()),
        "n_all_positive": int(((df["complete"]) & ~df["has_dash_zero"]).sum()),
        "n_with_dash_zero": int(df["has_dash_zero"].sum()),
        "cell_kinds": kind_counts,
        "n_labelled_cells": int(df["n_labelled"].sum()),
        "n_label_inconsistent": int(df["label_inconsistent_count"].sum()),
        "mean_mm_median": float(df["mean_mm"].median()),
        "mean_mm_p95": float(df["mean_mm"].quantile(0.95)),
        "mean_mm_max": float(df["mean_mm"].max()),
        "n_age": int(df["age"].notna().sum()),
        "n_age_from_sheet": int((df["age_source"] == "sheet").sum()),
        "n_age_from_prefix": int((df["age_source"] == "prefix").sum()),
        "n_sex": int(df["sex"].notna().sum()),
        "n_prefixed": int(df["age_prefix"].notna().sum()),
        "n_prefix_matches_sheet": int(df["prefix_age_matches_sheet"].sum()),
        "n_duplicate_keys": int(df["key"].duplicated(keep=False).sum()),
        "duplicate_keys": sorted(df.loc[df["key"].duplicated(keep=False), "key"].unique().tolist()),
    }


def read_demo_sheet(path: Path, sheet: str, excel_cfg: Dict[str, Any]) -> pd.DataFrame:
    """Name / age / sex rows of the low or normal sheet."""
    cols = excel_cfg["demo_columns"]
    records = []
    for ridx, row in enumerate(_rows(path, sheet)):
        raw_name = _cell(row, cols["name"])
        if ridx == 0 or raw_name is None or str(raw_name).strip() == "":
            continue
        name = normalize_excel_name(raw_name)
        records.append({
            "excel_row": ridx, "raw_name": str(raw_name).strip(), "key": name.key, "base": name.base,
            "dot": name.dot, "age_prefix": name.age_prefix,
            "age": parse_age(_cell(row, cols["age"])), "sex": parse_sex(_cell(row, cols["sex"])),
            "sheet": sheet,
        })
    df = pd.DataFrame(records)
    df["age"] = df["age"].astype(float)
    df["sex"] = df["sex"].astype(object).where(df["sex"].notna(), None)
    return df


def read_calibration(path: Path, excel_cfg: Dict[str, Any]) -> pd.DataFrame:
    """Long table: image, session (1|2), tooth, mm, parse_kind for the intra-observer file."""
    ccfg = excel_cfg["calibration"]
    teeth = excel_cfg["teeth"]
    records = []
    for session, sheet in ((1, ccfg["first_sheet"]), (2, ccfg["second_sheet"])):
        for ridx, row in enumerate(_rows(path, sheet)):
            if ridx < int(ccfg["header_rows"]):
                continue
            raw_name = _cell(row, ccfg["name_column"])
            if raw_name is None or str(raw_name).strip() == "":
                continue
            for i, col in enumerate(ccfg["teeth_columns"]):
                cell = parse_measurement_cell(_cell(row, col))
                records.append({
                    "image": str(raw_name).strip(), "key": normalize_excel_name(raw_name).key,
                    "session": session, "tooth": teeth[i], "tooth_index": i + 1,
                    "mm": cell.value_mm, "parse_kind": cell.parse_kind,
                })
    return pd.DataFrame(records)
