"""Expert evaluation forms (spec: docs/Uzman_degerlendirme_protokolu.md + task §Aşama 4).

Layout: sheet ``Degerlendirme``, header on the 4th row (0-based 3), data from the 5th,
165 rows: ``Sıra | Görüntü ID | Ölçek (pixel/mm) | 13 | 12 | 11 | 21 | 22 | 23 |
Etiyoloji sınıfı | İkinci aday | Güven | Not``. Real forms are messy; every anomaly is
flagged in ``form_flags`` and the row is kept — nothing is dropped silently.
"""
from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np
import openpyxl
import pandas as pd

TEETH = [13, 12, 11, 21, 22, 23]
VALID_CLASSES = {"E1", "E2", "E3", "E4"}
_CLASS_RE = re.compile(r"^\s*E\s*([1-4])\s*$", re.IGNORECASE)
_COMBINED_RE = re.compile(r"^\s*E\s*([1-4])\s*[-–/]\s*E?\s*([1-4])\s*$", re.IGNORECASE)


def parse_class(raw: Any) -> tuple:
    """Return (primary, secondary_from_cell, flag)."""
    if raw is None or (isinstance(raw, float) and math.isnan(raw)) or str(raw).strip() == "":
        return None, None, "missing_class"
    s = str(raw).strip()
    m = _CLASS_RE.match(s)
    if m:
        return f"E{m.group(1)}", None, None
    m = _COMBINED_RE.match(s)
    if m:
        return f"E{m.group(1)}", f"E{m.group(2)}", "combined_class_in_cell"
    return None, None, "invalid_class"


def parse_number(raw: Any) -> tuple:
    """Numeric cell for mm or scale: (value or NaN, flag)."""
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        return math.nan, "missing"
    if isinstance(raw, (int, float, np.integer, np.floating)) and not isinstance(raw, bool):
        v = float(raw)
    else:
        try:
            v = float(str(raw).strip().replace(",", "."))
        except ValueError:
            return math.nan, "unparsed"
    if math.isnan(v):
        return math.nan, "missing"
    if v < 0:
        return math.nan, "negative"
    return v, None


def read_form(path: Path, expert_cfg: Dict[str, Any]) -> pd.DataFrame:
    """One record per form row; anomalies collected in ``form_flags`` (comma-joined)."""
    sheet = expert_cfg.get("form_sheet", "Degerlendirme")
    header_row = int(expert_cfg.get("form_header_row", 3))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"sheet {sheet!r} not in {path.name}: {wb.sheetnames}")
    rows = list(wb[sheet].iter_rows(values_only=True))
    records: List[Dict[str, Any]] = []
    for ridx, row in enumerate(rows):
        if ridx <= header_row:
            continue
        row = list(row) + [None] * (13 - len(row))
        sira, gid, scale_raw = row[0], row[1], row[2]
        if gid is None and all(c is None or str(c).strip() == "" for c in row[:13]):
            continue  # fully empty trailing row
        flags: List[str] = []
        gid_s = str(gid).strip() if gid is not None else ""
        if not gid_s:
            flags.append("missing_image_id")
        rec: Dict[str, Any] = {"form_row": ridx, "sira": sira, "goruntu_id": gid_s}
        scale, f = parse_number(scale_raw)
        rec["scale_px_per_mm"] = scale
        if f:
            flags.append(f"scale_{f}")
        elif scale == 0:
            rec["scale_px_per_mm"] = math.nan
            flags.append("scale_zero")
        n_mm = 0
        for i, tooth in enumerate(TEETH):
            v, f = parse_number(row[3 + i])
            rec[f"mm_{tooth}"] = v
            rec[f"mm_{i + 1}"] = v
            if f == "missing":
                flags.append(f"mm_{tooth}_missing")
            elif f:
                flags.append(f"mm_{tooth}_{f}")
            else:
                n_mm += 1
        rec["n_mm_filled"] = n_mm
        vals = [rec[f"mm_{t}"] for t in TEETH]
        rec["mean_mm"] = float(np.nanmean(vals)) if n_mm else math.nan
        rec["mean_mm_complete"] = float(np.mean(vals)) if n_mm == 6 else math.nan
        primary, sec_cell, f = parse_class(row[9])
        if f:
            flags.append(f)
        sec, _, f2 = parse_class(row[10])
        rec["class_primary"] = primary
        rec["class_secondary"] = sec if sec else sec_cell
        if f2 == "invalid_class":
            flags.append("invalid_secondary")
        conf, f = parse_number(row[11])
        if f:
            flags.append("confidence_missing" if f == "missing" else f"confidence_{f}")
            conf = math.nan
        elif not (1 <= conf <= 5):
            flags.append("confidence_out_of_range")
            conf = math.nan
        rec["confidence"] = conf
        rec["note"] = str(row[12]).strip() if row[12] is not None else ""
        rec["row_empty"] = not gid_s or (primary is None and n_mm == 0)
        if rec["row_empty"] and gid_s:
            flags.append("row_empty")
        rec["form_flags"] = ",".join(flags)
        records.append(rec)
    df = pd.DataFrame(records)
    df.attrs["path"] = str(path)
    return df


def join_key(form: pd.DataFrame, key: pd.DataFrame) -> pd.DataFrame:
    """Attach the original image and the repeat flag from ANAHTAR_arastirmaci.csv."""
    k = key.rename(columns={"orijinal": "image", "tekrar": "is_repeat"})[["goruntu_id", "sira", "image", "is_repeat"]].copy()
    k["is_repeat"] = k["is_repeat"].astype(int).astype(bool)
    out = form.merge(k.drop(columns=["sira"]), on="goruntu_id", how="left", validate="one_to_one")
    out["key_missing"] = out["image"].isna()
    out["form_flags"] = np.where(out["key_missing"], np.where(out["form_flags"] == "", "id_not_in_key", out["form_flags"] + ",id_not_in_key"), out["form_flags"])
    return out


def split_repeats(joined: pd.DataFrame) -> tuple:
    """(first evaluations, repeat pairs) — the repeat pair table has ``_first`` / ``_repeat`` columns."""
    first = joined[~joined["is_repeat"].fillna(False).astype(bool)].copy()
    rep = joined[joined["is_repeat"].fillna(False).astype(bool)].copy()
    pairs = rep.merge(first, on="image", suffixes=("_repeat", "_first"), how="inner")
    return first, pairs


def form_qc_summary(form: pd.DataFrame) -> Dict[str, Any]:
    flags = pd.Series([f for s in form["form_flags"] for f in str(s).split(",") if f])
    return {
        "n_rows": int(len(form)), "n_rows_empty": int(form["row_empty"].sum()),
        "n_class_missing": int(form["class_primary"].isna().sum()),
        "n_scale_missing": int(form["scale_px_per_mm"].isna().sum()),
        "n_mm_complete": int((form["n_mm_filled"] == 6).sum()),
        "n_confidence_missing": int(form["confidence"].isna().sum()),
        "flag_counts": flags.value_counts().to_dict(),
    }
