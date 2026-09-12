"""Excel parsing rules from spec §4.2 / audit §4 (test 11), on a synthetic workbook."""
import math

import openpyxl
import pytest

from gsv4.io.excel_parser import (
    parse_label,
    parse_measurement_cell,
    parse_summary,
    read_calibration,
    read_demo_sheet,
    read_high_sheet,
)

EXCEL_CFG = {
    "high_sheet": "Yüksek Gülme Hattı",
    "low_sheet": "Düşük Gülme Hattı",
    "normal_sheet": "Normal Gülme Hattı",
    "high_header_row": 1,
    "teeth": [13, 12, 11, 21, 22, 23],
    "high_columns": {
        "name": 0,
        "teeth": [1, 2, 3, 4, 5, 6],
        "etiology": [8, 9, 10, 11, 12, 13],
        "treatment": [14, 15, 16, 17, 18, 19],
        "age": 23,
        "sex": 24,
    },
    "demo_columns": {"name": 0, "age": 2, "sex": 3},
    "calibration": {
        "first_sheet": "İlk Ölçümler",
        "second_sheet": "İkinci Ölçümler",
        "header_rows": 3,
        "name_column": 0,
        "teeth_columns": [2, 3, 4, 5, 6, 7],
    },
}


@pytest.mark.parametrize(
    "raw,value,kind",
    [
        (3207, 3.207, "div1000"),
        (13768, 13.768, "div1000"),
        ("*0.718", 0.718, "star_sub1mm"),
        ("*0,718", 0.718, "star_sub1mm"),
        ("-", 0.0, "dash_zero"),
        ("-(mesafe yok)", 0.0, "dash_zero"),
        ("               -", 0.0, "dash_zero"),
        ("3.00", 3.0, "plain_mm"),
        ("3.44", 3.44, "plain_mm"),
        (3, 3.0, "plain_mm"),
        (4, 4.0, "plain_mm"),
        (3.5, 3.5, "plain_mm"),
        (-684, math.nan, "negative"),
        (712, 0.712, "ambiguous_100_999"),
        (None, math.nan, "missing"),
        ("", math.nan, "missing"),
        ("abc", math.nan, "unparsed"),
        (42, math.nan, "unparsed"),  # 10..99 has no defined meaning
    ],
)
def test_parse_measurement_cell(raw, value, kind):
    cell = parse_measurement_cell(raw)
    assert cell.parse_kind == kind
    if math.isnan(value):
        assert math.isnan(cell.value_mm)
    else:
        assert cell.value_mm == pytest.approx(value)


def test_parse_label():
    assert parse_label(" E1-E2 ") == "E1-E2"
    assert parse_label("e2-e3") == "E2-E3"
    assert parse_label("E1–E2") == "E1-E2"  # en dash
    assert parse_label("-") == "-"
    assert parse_label(None) is None
    assert parse_label("T1") == "T1"


def _build_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = EXCEL_CFG["high_sheet"]
    ws.append([None, "Soldan sağa 6 dişin numarası"])
    header = ["image numarası", 13, 12, 11, 21, 22, 23, None, 1, 2, 3, 4, 5, 6, 1, 2, 3, 4, 5, 6]
    ws.append(header + [None, None, None, "YAŞ", "CİNSİYET"])
    rows = [
        # name, six cells, blank, six E labels, six T labels, 3 blanks, age, sex
        ["IMG_0001", 2207, 2392, "*0.718", "-", 5000, 712, None,
         "E1", "E1", "E1", "-", "E2-E3", "E1", "T1", "T1", "T1", "-", "T2-T3", "T1", None, None, None, 25, "K"],
        ["IMG_0002.", 4000, 4000, 4000, 4000, 4000, 4000, None,
         "E2-E3", "E2-E3", "E2-E3", "E2-E3", "E2-E3", "E2-E3", "T2-T3"] + ["T2-T3"] * 5 + [None, None, None, "E: Erkek", "E"],
        ["25-IMG_0003", 1500, 1500, 1500, -684, 1500, 1500, None,
         "E1", "E1", "E1", "E1", "E1", "E1", "T1", "T1", "T1", "T1", "T1", "T1", None, None, None, None, None],
        # label inconsistent: 667 read as 0.667 (E1) but labelled E4
        ["IMG_0004", 9000, 9000, 9000, 9000, 9000, 667, None,
         "E4", "E4", "E4", "E4", "E4", "E4", "T4"] + ["T4"] * 5 + [None, None, None, 40, "K"],
        [None] * 25,
    ]
    for r in rows:
        ws.append(r)
    ws_low = wb.create_sheet(EXCEL_CFG["low_sheet"])
    ws_low.append([None, None, "YAŞ", "CİNSİYET"])
    ws_low.append(["IMG_9545", None, 37, "K"])
    ws_low.append(["IMG_5669.", None, 30, "E"])
    ws_low.append(["IMG_5670", None, None, None])
    ws_n = wb.create_sheet(EXCEL_CFG["normal_sheet"])
    ws_n.append([None, None, "YAŞ", "CİNSİYET"])
    ws_n.append(["IMG_78701.jpg", None, 25, "K"])
    wb.save(path)


def test_read_high_sheet(tmp_path):
    path = tmp_path / "m.xlsx"
    _build_workbook(path)
    df = read_high_sheet(path, EXCEL_CFG)
    assert len(df) == 4  # trailing empty row dropped
    r1 = df.iloc[0]
    assert r1["key"] == "img0001"
    assert r1["mm_1"] == pytest.approx(2.207)
    assert r1["mm_3"] == pytest.approx(0.718)
    assert r1["mm_4"] == 0.0 and r1["kind_4"] == "dash_zero"
    assert r1["kind_6"] == "ambiguous_100_999"
    assert r1["n_valid"] == 6 and bool(r1["complete"]) is True
    assert bool(r1["has_dash_zero"]) is True
    assert r1["age"] == 25 and r1["sex"] == "F"
    assert r1["label_4"] == "-" and r1["expected_label_4"] == "-"
    assert r1["label_inconsistent_count"] == 0
    # mean over six teeth (dash counts as 0)
    assert r1["mean_mm"] == pytest.approx((2.207 + 2.392 + 0.718 + 0 + 5.0 + 0.712) / 6)

    r2 = df.iloc[1]
    assert r2["key"] == "img0002~" and bool(r2["dot"]) is True
    assert math.isnan(r2["age"]) and r2["age_raw"] == "E: Erkek"
    assert r2["sex"] == "M"

    r3 = df.iloc[2]
    assert r3["key"] == "25-img0003" and r3["age_prefix"] == 25
    assert r3["age"] == 25 and r3["age_source"] == "prefix"
    assert r3["kind_4"] == "negative" and math.isnan(r3["mm_4"])
    assert r3["n_valid"] == 5 and bool(r3["complete"]) is False

    r4 = df.iloc[3]
    assert r4["label_inconsistent_count"] == 1
    assert r4["expected_label_6"] == "E1" and r4["label_6"] == "E4"

    s = parse_summary(df)
    assert s["n_rows"] == 4
    assert s["n_complete"] == 3
    assert s["cell_kinds"]["dash_zero"] == 1
    assert s["cell_kinds"]["star_sub1mm"] == 1
    assert s["cell_kinds"]["negative"] == 1
    assert s["n_labelled_cells"] == 24
    assert s["n_label_inconsistent"] == 1


def test_read_demo_sheet(tmp_path):
    path = tmp_path / "m.xlsx"
    _build_workbook(path)
    low = read_demo_sheet(path, EXCEL_CFG["low_sheet"], EXCEL_CFG)
    assert list(low["key"]) == ["img9545", "img5669~", "img5670"]
    assert list(low["sex"]) == ["F", "M", None]
    normal = read_demo_sheet(path, EXCEL_CFG["normal_sheet"], EXCEL_CFG)
    assert normal.iloc[0]["key"] == "img78701"


def test_read_calibration(tmp_path):
    wb = openpyxl.Workbook()
    for title, base in ((EXCEL_CFG["calibration"]["first_sheet"], 3750), (EXCEL_CFG["calibration"]["second_sheet"], 3800)):
        ws = wb.create_sheet(title)
        ws.append(["RESİM", None, 1, 2, 3, 4, 5, 6])
        ws.append([None, None, "Soldan sağa"])
        ws.append([None, None, 13, 12, 11, 21, 22, 23])
        ws.append(["IMG_8607", None, base, base + 1, base + 2, base + 3, base + 4, base + 5])
        ws.append(["IMG_8609", None, "*0.500", 4598, 3736, 3710, 4890, "-"])
    del wb["Sheet"]
    path = tmp_path / "c.xlsx"
    wb.save(path)
    long = read_calibration(path, EXCEL_CFG)
    assert len(long) == 2 * 2 * 6
    first = long[(long.session == 1) & (long.image == "IMG_8607") & (long.tooth == 13)]
    assert first.iloc[0]["mm"] == pytest.approx(3.750)
    assert long[(long.session == 2) & (long.image == "IMG_8609") & (long.tooth == 13)].iloc[0]["mm"] == pytest.approx(0.5)
    assert long[(long.image == "IMG_8609") & (long.tooth == 23)].iloc[0]["mm"] == 0.0
